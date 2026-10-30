from pathlib import Path
from typing import Iterable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import settings

DOC_EXTS = {".pdf", ".md", ".txt", ".docx"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp"}


def load_document(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader
        docs = PyPDFLoader(str(path)).load()
    elif suffix == ".docx":
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        docs = [Document(page_content=text, metadata={"source": path.name})]
    elif suffix in {".md", ".txt"}:
        docs = [Document(page_content=path.read_text(encoding="utf-8", errors="replace"), metadata={"source": path.name})]
    else:
        raise ValueError(f"Unsupported document type: {suffix}")
    for doc in docs:
        doc.metadata.update(source=path.name, file_type="document")
    return docs


def ocr_image(path: Path) -> str:
    from paddleocr import PaddleOCR
    ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
    result = ocr.ocr(str(path), cls=True)
    lines = []
    for page in result or []:
        for item in page or []:
            if item and len(item) > 1:
                lines.append(str(item[1][0]))
    return "\n".join(lines).strip()


def split_documents(docs: Iterable[Document], size: int | None = None, overlap: int | None = None) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=size or settings.chunk_size, chunk_overlap=overlap or settings.chunk_overlap, add_start_index=True)
    return splitter.split_documents(list(docs))
